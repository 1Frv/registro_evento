import io
import logging
import shutil
import sqlite3
from datetime import datetime
from pathlib import Path

from django.conf import settings
from django.db import connection
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font

from .models import Asistente

log = logging.getLogger(__name__)


def respaldar():
    """Copia consistente de la base (API de respaldo de SQLite) + copia opcional a otra ubicación."""
    destino = Path(settings.RESPALDOS_DIR)
    destino.mkdir(exist_ok=True)
    archivo = destino / f"registro-{datetime.now():%Y%m%d-%H%M%S}.sqlite3"
    connection.ensure_connection()
    copia = sqlite3.connect(archivo)
    try:
        connection.connection.backup(copia)
    finally:
        copia.close()
    if settings.RESPALDO_EXTRA_DIR:
        try:
            Path(settings.RESPALDO_EXTRA_DIR).mkdir(parents=True, exist_ok=True)
            shutil.copy2(archivo, settings.RESPALDO_EXTRA_DIR)
        except OSError:
            log.warning("No se pudo copiar el respaldo a %s", settings.RESPALDO_EXTRA_DIR)
    for viejo in sorted(destino.glob("registro-*.sqlite3"))[:-200]:
        viejo.unlink()
    return archivo


def exportar_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Asistentes"
    cols = [("Folio", 8), ("Fecha y hora", 18), ("Nombre", 30), ("Institución / Empresa / Municipio", 36),
            ("Cargo", 22), ("Correo", 28), ("Teléfono", 16), ("Extensión", 12), ("Firma", 26)]
    ws.append([c for c, _ in cols])
    for i, (_, ancho) in enumerate(cols):
        ws.column_dimensions[chr(65 + i)].width = ancho
    for celda in ws[1]:
        celda.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for fila, a in enumerate(Asistente.objects.filter(anulado=False).order_by("id"), start=2):
        ws.append([a.id, timezone.localtime(a.creado).strftime("%d/%m/%Y %H:%M"),
                   a.nombre, a.institucion, a.cargo, a.correo, a.telefono, a.extension])
        ws.row_dimensions[fila].height = 52
        for celda in ws[fila]:
            celda.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(celda.value, str):
                celda.data_type = "s"  # texto literal: nunca se interpreta como fórmula
        img = XLImage(io.BytesIO(bytes(a.firma)))
        img.width, img.height = 150, 60
        ws.add_image(img, f"{chr(64 + len(cols))}{fila}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()